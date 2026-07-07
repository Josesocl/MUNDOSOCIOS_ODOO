"""Tests del generador de devengos. Ejecutar desde generador-devengos/:
    python3 -m pytest tests/ -q     (o)     python3 -m unittest discover tests
"""

import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import rut_utils
import generador_devengos as gd
import generador_cuota_social as gcs


class TestRut(unittest.TestCase):
    def test_normaliza_puntos_y_k(self):
        self.assertEqual(rut_utils.normalizar("5.080.884-k"), "5080884-K")
        self.assertEqual(rut_utils.normalizar(" 4674134k "), "4674134-K")

    def test_dv_correctos_del_manual(self):
        # RUTs reales del ejemplo del MANUAL CARGA DE DEVENGOS
        for rut in ["5031324-7", "5080884-K", "4674134-K", "3543203-5",
                    "4461179-1", "6551451-6"]:
            self.assertTrue(rut_utils.es_valido(rut), rut)

    def test_dv_incorrecto(self):
        self.assertFalse(rut_utils.es_valido("5031324-8"))
        self.assertFalse(rut_utils.es_valido("no-es-rut"))


class TestPeriodoYMonto(unittest.TestCase):
    def test_parse_periodo(self):
        self.assertEqual(gd.parse_periodo("07-2026"), (7, 2026))
        self.assertEqual(gd.parse_periodo("2026-07"), (7, 2026))
        self.assertEqual(gd.parse_periodo("7/2026"), (7, 2026))
        with self.assertRaises(ValueError):
            gd.parse_periodo("13-2026")

    def test_etiquetas(self):
        self.assertEqual(gd.etiqueta_periodo(6, 2024), "JUN 24")
        self.assertEqual(gd.fecha_contable(6, 2024), date(2024, 6, 1))

    def test_monto_redondeo_al_peso(self):
        # 3 UF a UF=$37.439,17 -> $112.317,51 -> $112.318
        self.assertEqual(gd.monto_clp(Decimal("3"), Decimal("37439.17")), 112318)
        # caso del manual: factor 2 con UF ~37.439 -> $74.878
        self.assertEqual(gd.monto_clp(Decimal("2"), Decimal("37439")), 74878)

    def test_normalizar_seguro(self):
        self.assertEqual(gd.normalizar_seguro("Plan Carreño"), "PLAN CARRENO")
        self.assertEqual(gd.normalizar_seguro("catastrófico"), "CATASTROFICO")
        with self.assertRaises(ValueError):
            gd.normalizar_seguro("SEGURO INEXISTENTE")


class TestMaestro(unittest.TestCase):
    def _maestro(self, contenido):
        d = tempfile.mkdtemp()
        ruta = Path(d) / "maestro.csv"
        ruta.write_text(contenido, encoding="utf-8")
        return ruta

    def test_maestro_valido(self):
        ruta = self._maestro(
            "rut;nombre;seguro;factor_uf;medio_pago\n"
            "5031324-7;CARLOS PRIETO;PLAN CARRENO;0,4;PAC\n"
            "5080884-k;SARA MUNOZ;Plan Carreño;0,4;DIRECTA\n"
            "4461179-1;RENE POBLETE;PLAN SOCIOS;3;PAC\n")
        filas, errores = gd.leer_maestro(ruta)
        self.assertEqual(errores, [])
        self.assertEqual(len(filas), 3)
        self.assertEqual(filas[1]["rut"], "5080884-K")  # K a mayúscula
        self.assertEqual(filas[1]["seguro"], "PLAN CARRENO")

    def test_maestro_detecta_errores(self):
        ruta = self._maestro(
            "rut;nombre;seguro;factor_uf\n"
            "5031324-8;DV MALO;PLAN SOCIOS;2\n"          # DV incorrecto
            "5031324-7;OK;PLAN SOCIOS;2\n"
            "5031324-7;DUPLICADO;PLAN SOCIOS;2\n"        # duplicado
            "4461179-1;FACTOR MALO;CATASTROFICO;0\n"     # factor <= 0
            "6551451-6;SEGURO MALO;VIDA ENTERA;1\n")     # seguro desconocido
        filas, errores = gd.leer_maestro(ruta)
        texto = "\n".join(errores)
        self.assertIn("dígito verificador incorrecto", texto)
        self.assertIn("aparece 2 veces", texto)
        self.assertIn("factor_uf debe ser > 0", texto)
        self.assertIn("Seguro desconocido", texto)

    def test_columna_faltante(self):
        ruta = self._maestro("rut;nombre\n1-9;X\n")
        _, errores = gd.leer_maestro(ruta)
        self.assertTrue(any("no tiene las columnas" in e for e in errores))


class TestArchivo(unittest.TestCase):
    def test_estructura_y_contrapartida(self):
        polizas = [
            {"rut": "5031324-7", "nombre": "A", "seguro": "PLAN CARRENO",
             "factor_uf": Decimal("0.4"), "medio_pago": "PAC"},
            {"rut": "5080884-K", "nombre": "B", "seguro": "PLAN CARRENO",
             "factor_uf": Decimal("0.2"), "medio_pago": ""},
        ]
        filas, total = gd.construir_filas(polizas, "PLAN CARRENO", 5, 2024,
                                          Decimal("37307.5"))
        self.assertEqual(len(filas), 3)  # 2 detalles + contrapartida
        det1, det2, contra = filas
        for fila in filas:
            self.assertEqual(len(fila), 22)  # columnas A–V
        # Detalle: T, glosa, fecha, correlativo, UN 001, cuenta CxC, tipo doc
        self.assertEqual(det1[0], "T")
        self.assertEqual(det1[2], "DEVENGO P CARREÑO MAY 24")
        self.assertEqual(det1[3], date(2024, 5, 1))
        self.assertEqual([det1[6], det2[6]], [1, 2])
        self.assertEqual(det1[7], "001")
        self.assertEqual(det1[8], "5031324-7 P CARR MAY-24")  # detalle con guion
        self.assertEqual(det1[11], "1130005")
        self.assertEqual(det1[12], "")           # CC solo en última fila
        self.assertEqual(det1[13], 14923)        # 0,4 × 37.307,5 = 14.923
        self.assertEqual(det1[14], 0)
        self.assertEqual(det1[15], "PCARR")
        self.assertEqual([det1[18], det1[19]], ["99999", "500"])  # Conceptos 1-2
        # Contrapartida: cuenta ingreso, CC MS, HABER = suma, sin tipo doc
        self.assertEqual(contra[6], 3)
        self.assertEqual(contra[9], "")
        self.assertEqual(contra[11], "3310004")
        self.assertEqual(contra[12], "MS")
        self.assertEqual(contra[13], 0)
        self.assertEqual(contra[14], total)
        self.assertEqual(contra[15], "")
        self.assertEqual(total, 14923 + 7462)    # 0,2 × 37.307,5 = 7.461,5 -> 7.462

    def test_xlsx_legible(self):
        from openpyxl import load_workbook
        polizas = [{"rut": "5031324-7", "nombre": "A", "seguro": "PLAN SOCIOS",
                    "factor_uf": Decimal("3"), "medio_pago": "PAC"}]
        filas, total = gd.construir_filas(polizas, "PLAN SOCIOS", 6, 2024,
                                          Decimal("37439"))
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "out.xlsx"
            gd.escribir_xlsx(filas, ruta)
            ws = load_workbook(ruta).active
            self.assertEqual(ws.max_row, 3)      # encabezado + detalle + contra
            self.assertEqual(ws.max_column, 22)
            self.assertEqual(ws.cell(1, 1).value, "Tipo de comprobante")
            self.assertEqual(ws.cell(2, 10).value, "5031324-7")
            self.assertEqual(ws.cell(2, 14).value, 112317)   # 3 × 37.439
            self.assertEqual(ws.cell(3, 15).value, total)


class TestCli(unittest.TestCase):
    def test_flujo_completo(self):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            (d / "maestro.csv").write_text(
                "rut,nombre,seguro,factor_uf,medio_pago\n"
                "5031324-7,CARLOS PRIETO,PLAN CARRENO,0.4,PAC\n"
                "4461179-1,RENE POBLETE,PLAN SOCIOS,3,PAC\n", encoding="utf-8")
            (d / "clientes.csv").write_text("rut\n5031324-7\n", encoding="utf-8")
            rc = gd.main(["--maestro", str(d / "maestro.csv"),
                          "--periodo", "05-2024", "--uf", "37307.5",
                          "--clientes", str(d / "clientes.csv"),
                          "--salida", str(d / "out")])
            self.assertEqual(rc, 0)
            generados = sorted(p.name for p in (d / "out").glob("*.xlsx"))
            self.assertEqual(generados, ["DEVENGO PLAN CARREÑO MAY-24.xlsx",
                                         "DEVENGO PLAN SOCIOS MAY-24.xlsx"])

    def test_maestro_con_errores_no_genera(self):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            (d / "maestro.csv").write_text(
                "rut,nombre,seguro,factor_uf\n"
                "5031324-8,DV MALO,PLAN SOCIOS,2\n", encoding="utf-8")
            rc = gd.main(["--maestro", str(d / "maestro.csv"),
                          "--periodo", "05-2024", "--uf", "37307.5",
                          "--salida", str(d / "out")])
            self.assertEqual(rc, 1)
            self.assertEqual(list((d / "out").glob("*.xlsx")) if (d / "out").exists() else [], [])


class TestCuotaSocial(unittest.TestCase):
    def test_unidades(self):
        self.assertEqual(gcs.unidades_uf("PERSONA", ""), Decimal(1))
        self.assertEqual(gcs.unidades_uf("EMPRESA", ""), Decimal(3))
        self.assertEqual(gcs.unidades_uf("EMPRESA", "2"), Decimal(3))
        self.assertEqual(gcs.unidades_uf("EMPRESA", "3"), Decimal(3))
        self.assertEqual(gcs.unidades_uf("EMPRESA", "5"), Decimal(5))

    def test_montos_reales_2026(self):
        # UF del devengo ENE-26: base empresa $119.195 (3 UF), 4 UF $158.927,
        # persona $39.732 — verificados contra los archivos reales
        uf = Decimal("39731.77")
        socios = [
            {"rut": "5031324-7", "nombre": "A", "tipo": "EMPRESA",
             "camara": "O'HIGGINS", "miembros": "", "monto_manual": None},
            {"rut": "4461179-1", "nombre": "B", "tipo": "EMPRESA",
             "camara": "SANTIAGO", "miembros": "4", "monto_manual": None},
        ]
        filas, total = gcs.construir_filas(socios, "EMPRESA", 2026, uf)
        det1, det2, contra = filas
        self.assertEqual(det1[13], 119195)
        self.assertEqual(det2[13], 158927)
        self.assertEqual(det1[2], "DEVENGO CUOTA EMPRESA 2026")
        self.assertEqual(det1[8], "5031324-7 CE 26 O'HIGGINS")
        self.assertEqual(det1[11], "1150001")
        self.assertEqual(det1[15], "CSEMP")
        self.assertEqual([det1[18], det1[19]], ["", ""])   # sin conceptos
        self.assertEqual(contra[11], "3210002")
        self.assertEqual(contra[12], "ADM")
        self.assertEqual(contra[14], total)

    def test_persona_cuentas_cruzadas_y_override(self):
        uf = Decimal("39731.77")
        socios = [
            {"rut": "5031324-7", "nombre": "A", "tipo": "PERSONA",
             "camara": "MAULE", "miembros": "", "monto_manual": None},
            {"rut": "4461179-1", "nombre": "B", "tipo": "PERSONA",
             "camara": "ARICA", "miembros": "", "monto_manual": 39587},
        ]
        filas, total = gcs.construir_filas(socios, "PERSONA", 2026, uf)
        det1, det2, contra = filas
        self.assertEqual(det1[13], 39732)          # 1 UF
        self.assertEqual(det2[13], 39587)          # excepción manual
        self.assertEqual(det1[8], "5031324-7 CP 26 MAULE")
        self.assertEqual(det1[11], "1150002")      # cuentas cruzadas reales
        self.assertEqual(contra[11], "3210001")
        self.assertEqual(det1[15], "CSPER")
        self.assertEqual(total, 39732 + 39587)

    def test_cli(self):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            (d / "maestro.csv").write_text(
                "rut;nombre;tipo_socio;camara;miembros;monto_clp\n"
                "5031324-7;EMPRESA UNO;EMPRESA;SANTIAGO;;\n"
                "4461179-1;PERSONA UNO;PERSONA;MAULE;;\n", encoding="utf-8")
            rc = gcs.main(["--maestro", str(d / "maestro.csv"), "--anio", "2026",
                           "--uf", "39731.77", "--salida", str(d / "out")])
            self.assertEqual(rc, 0)
            generados = sorted(p.name for p in (d / "out").glob("*.xlsx"))
            self.assertEqual(generados, ["DEVENGO CUOTA SOCIAL EMPRESA ENE-26.xlsx",
                                         "DEVENGO CUOTA SOCIAL PERSONA ENE-26.xlsx"])


class TestUfPorSeguro(unittest.TestCase):
    def test_uf_especifica(self):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            (d / "maestro.csv").write_text(
                "rut,nombre,seguro,factor_uf\n"
                "5031324-7,A,PLAN CARRENO,0.2\n"
                "4461179-1,B,CATASTROFICO,1\n", encoding="utf-8")
            rc = gd.main(["--maestro", str(d / "maestro.csv"),
                          "--periodo", "07-2026", "--uf", "40845",
                          "--uf-seguro", "CATASTROFICO=40763",
                          "--salida", str(d / "out")])
            self.assertEqual(rc, 0)
            from openpyxl import load_workbook
            ws = load_workbook(d / "out" / "DEVENGO PLAN CARREÑO JUL-26.xlsx").active
            self.assertEqual(ws.cell(2, 14).value, 8169)    # 0,2 × 40.845
            ws = load_workbook(d / "out" / "DEVENGO CATASTRÓFICO JUL-26.xlsx").active
            self.assertEqual(ws.cell(2, 14).value, 40763)   # UF específica


if __name__ == "__main__":
    unittest.main()
