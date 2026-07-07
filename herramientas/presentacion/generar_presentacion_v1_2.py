#!/usr/bin/env python3
"""Genera Presentacion_Ejecutiva_MS_JR_V1.2.pptx.

Informe ejecutivo de la asesoría (COT-2026-MS-001) para la Gerencia
General de MundoSocios, estructurado en: lo que está listo, lo que falta
por terminar, lo que se va a automatizar y lo que no (etapa abierta),
trabajo con el equipo MS ampliado, y el puente hacia Odoo (nada se
pierde). Sin tablas de rendición: el trabajo se hizo.

Uso:  python3 generar_presentacion_v1_2.py [salida.pptx]
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

AZUL = RGBColor(0x1F, 0x4E, 0x79)
AZUL_CLARO = RGBColor(0xDC, 0xE6, 0xF2)
GRIS = RGBColor(0x59, 0x59, 0x59)
GRIS_CLARO = RGBColor(0xF2, 0xF2, 0xF2)
VERDE = RGBColor(0x2E, 0x7D, 0x32)
AMBAR = RGBColor(0xB5, 0x6A, 0x00)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
NEGRO = RGBColor(0x21, 0x21, 0x21)

ANCHO, ALTO = Inches(13.333), Inches(7.5)


def _caja(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.text_frame.word_wrap = True
    return tb


def _p(tf, texto, size=14, bold=False, color=NEGRO, first=False, align=PP_ALIGN.LEFT,
       space_after=6, bullet=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    r = p.add_run()
    r.text = ("•  " if bullet else "") + texto
    f = r.font
    f.size, f.bold, f.color.rgb, f.name = Pt(size), bold, color, "Calibri"
    p.alignment = align
    p.space_after = Pt(space_after)
    return p


def _rect(slide, x, y, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def slide_base(prs, titulo, subtitulo=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, ANCHO, Inches(0.16), AZUL)
    tb = _caja(s, Inches(0.5), Inches(0.32), Inches(12.3), Inches(0.75))
    _p(tb.text_frame, titulo, size=26, bold=True, color=AZUL, first=True)
    if subtitulo:
        _p(tb.text_frame, subtitulo, size=13, color=GRIS)
    pie = _caja(s, Inches(0.5), Inches(7.08), Inches(12.3), Inches(0.35))
    _p(pie.text_frame, "JR Jottar Consultoría · IB Solución Ltda. · Julio 2026",
       size=9, color=GRIS, first=True)
    return s


def bloque(slide, x, y, w, h, titulo, items, color_titulo=AZUL, size=13, fill=None):
    if fill:
        _rect(slide, x, y, w, h, fill)
    tb = _caja(slide, x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), h - Inches(0.2))
    _p(tb.text_frame, titulo, size=16, bold=True, color=color_titulo, first=True)
    for it in items:
        _p(tb.text_frame, it, size=size, bullet=True)
    return tb


def construir(salida: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = ANCHO, ALTO

    # ---- 1. Portada
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, ANCHO, ALTO, AZUL)
    tb = _caja(s, Inches(1), Inches(2.2), Inches(11.3), Inches(3.2))
    _p(tb.text_frame, "Asesoría Procesos Financiero-Contables", size=40, bold=True,
       color=BLANCO, first=True)
    _p(tb.text_frame, "MundoSocios — Cámara Chilena de la Construcción", size=22, color=AZUL_CLARO)
    _p(tb.text_frame, "", size=10)
    _p(tb.text_frame, "Informe Ejecutivo para la Gerencia General", size=18, color=BLANCO)
    _p(tb.text_frame, "Qué está listo · Qué falta · Qué se automatiza · Camino a Odoo",
       size=14, color=AZUL_CLARO)
    tb = _caja(s, Inches(1), Inches(6.3), Inches(11.3), Inches(0.6))
    _p(tb.text_frame, "JR Jottar Consultoría (IB Solución) · Julio 2026 · V1.2",
       size=12, color=AZUL_CLARO, first=True)

    # ---- 2. Contexto y objetivo
    s = slide_base(prs, "Contexto y objetivo")
    kpis = [("~1.955", "socios titulares"), ("~4.800", "personas con familias"),
            ("~1.500", "notas de cobro/mes"), ("4+", "sistemas no integrados")]
    x = Inches(0.5)
    for valor, etiqueta in kpis:
        _rect(s, x, Inches(1.35), Inches(2.95), Inches(0.95), GRIS_CLARO)
        tb = _caja(s, x, Inches(1.4), Inches(2.95), Inches(0.9))
        _p(tb.text_frame, valor, size=22, bold=True, color=AZUL, first=True,
           align=PP_ALIGN.CENTER, space_after=0)
        _p(tb.text_frame, etiqueta, size=11, color=GRIS, align=PP_ALIGN.CENTER)
        x += Inches(3.12)
    tb = _caja(s, Inches(0.5), Inches(2.7), Inches(12.3), Inches(1.6))
    _p(tb.text_frame,
       "MundoSocios opera hoy repartido entre Zoho, Manager+, WordPress, planillas y correo. "
       "La migración a Odoo Enterprise consolidará todo en un único sistema; mientras llega "
       "(no antes de noviembre), el área financiero-contable necesita alivio operativo ahora.",
       size=15, first=True)
    tb = _caja(s, Inches(0.5), Inches(4.2), Inches(12.3), Inches(2.4))
    _p(tb.text_frame, "El trabajo tiene dos frentes:", size=16, bold=True, color=AZUL, first=True)
    _p(tb.text_frame, "Aliviar los flujos de trabajo de hoy con automatizaciones que reduzcan "
       "horas manuales y errores.", size=15, bullet=True)
    _p(tb.text_frame, "Preparar a las personas y los datos para que el cambio a un único "
       "sistema llegue con el equipo convencido y la información limpia.", size=15, bullet=True)

    # ---- 3. Lo que está listo
    s = slide_base(prs, "Lo que está listo", "Construido, verificado y en manos del equipo")
    bloque(s, Inches(0.5), Inches(1.5), Inches(6.0), Inches(2.6),
           "Diagnóstico y documentación", [
               "Procesos financiero-contables levantados de punta a punta, con sus dolores y riesgos",
               "Blueprints de compras, recaudación y experiencia del socio (AS-IS → diseño Odoo)",
               "Los 4 entregables comprometidos, más guía de entrevistas y requerimientos Odoo",
           ])
    bloque(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(2.6),
           "Herramientas operativas", [
               "5 planillas de trabajo verificadas: ficha proveedor, OC, control de nómina/TXT, UF, conciliación",
               "Validador SII con proxy en la nube ya desplegado",
               "Generadores de devengos (seguros y cuota social): montos idénticos a los archivos reales",
           ])
    bloque(s, Inches(0.5), Inches(4.2), Inches(12.4), Inches(2.3),
           "Cobro recurrente (cuotas y seguros)", [
               "Cruzador de pagos: lee la cartola y Transbank, y arma el borrador de preconciliación solo",
               "Reglas 2026 de cuota social y de UF por seguro incorporadas a las herramientas",
               "Política de morosidad y pagos parciales redactada (en borrador para decisión)",
           ], color_titulo=VERDE)

    # ---- 4. Lo que falta por terminar
    s = slide_base(prs, "Lo que falta por terminar",
                   "Nada requiere construcción mayor: son cierres, decisiones y adopción")
    bloque(s, Inches(0.5), Inches(1.5), Inches(6.0), Inches(4.9),
           "Del lado del equipo MS", [
               "Entrevistas de cierre: caja chica, ajustes de cuentas por pagar, morosidad histórica",
               "Visto bueno de Patricio a la política de aprobaciones y al checklist de proveedores",
               "Decidir la política de morosidad y pagos parciales (borrador listo)",
               "Contratar el servicio de consulta SII (API Gateway, ~$10.000/mes) para que la validación deje de ser manual",
               "Pedir al banco la rendición PAC por convenio (para distribuir la recaudación por socio)",
           ])
    bloque(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(4.9),
           "Del lado del consultor", [
               "Consolidar el maestro único de socios y pólizas desde los mantenedores actuales",
               "Correr agosto en paralelo: generadores y cruzador junto al proceso manual, cuadrando resultados",
               "Entrenar al equipo en el uso semanal de las herramientas",
               "Calibrar el cruzador contra la preconciliación real de junio",
           ])

    # ---- 5. Lo que se va a automatizar
    s = slide_base(prs, "Lo que se va a automatizar",
                   "Para aliviar los flujos de trabajo entre hoy y la llegada de Odoo")
    bloque(s, Inches(0.5), Inches(1.5), Inches(6.0), Inches(4.6),
           "Devengo y cobro", [
               "Devengos mensuales de los 4 seguros generados desde un maestro único (se acaba el copiar el Excel del mes anterior)",
               "Devengo anual de cuota social generado con las reglas 2026",
               "Validación de RUTs y clientes ANTES de importar (se acaba el ciclo \"Cliente no existe\")",
               "Correo de incorporación con UF del día y monto calculado automáticamente",
           ])
    bloque(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(4.6),
           "Cruce, cobranza y conciliación", [
               "Cruce semanal de pagos: borrador de preconciliación automático con propuesta de socio",
               "Actualización masiva de estados en Zoho desde el cruce (\"cuota al día\" / \"moroso\")",
               "Paquete diario estándar para Addval, con control del plazo de 48 horas",
               "Correos de cobro y recordatorio por Zoho Campaigns, por fin con métricas",
           ])
    tb = _caja(s, Inches(0.5), Inches(6.15), Inches(12.3), Inches(0.8))
    _p(tb.text_frame, "Alivio estimado del paquete completo: 50–60 horas al mes en el área.",
       size=16, bold=True, color=VERDE, first=True)

    # ---- 6. Lo que NO se va a automatizar
    s = slide_base(prs, "Lo que no se va a automatizar (por ahora)",
                   "Esta etapa aún no está cerrada: la lista puede ajustarse con las decisiones pendientes")
    tb = _caja(s, Inches(0.5), Inches(1.6), Inches(12.3), Inches(4.4))
    _p(tb.text_frame, "El cargo automático PAC/PAT y Webpay de punta a punta — requiere integrar "
       "una pasarela de pagos; es exactamente el desarrollo que se hará en Odoo y no vale la pena "
       "pagarlo dos veces por tres meses.", size=14, bullet=True, first=True)
    _p(tb.text_frame, "El registro contable de los pagos en Manager+ — es la función contratada a "
       "Addval; primero se formaliza el ciclo diario y su plazo de 48 horas, y solo si no se cumple "
       "se evalúa internalizarlo.", size=14, bullet=True)
    _p(tb.text_frame, "La conciliación de convenios PAC (un documento contra muchos abonos) — "
       "seguirá manual-asistida; Odoo la trae resuelta de forma nativa.", size=14, bullet=True)
    _p(tb.text_frame, "La separación de fondos propios y de terceros — en el puente queda como "
       "reporte de control mensual; la separación estructural e inviolable la da Odoo con cuentas "
       "y diarios por producto.", size=14, bullet=True)
    _p(tb.text_frame, "La identificación automática de cada venta Transbank con su socio — los "
       "reportes del banco no traen RUT; se resuelve al guardar la orden de pago, parte del diseño "
       "de cobro en Odoo.", size=14, bullet=True)
    tb = _caja(s, Inches(0.5), Inches(6.05), Inches(12.3), Inches(0.9))
    _p(tb.text_frame, "Criterio: no construir nada desechable. Lo que Odoo resuelve nativo, se "
       "espera; lo que duele hoy y sobrevive a la migración, se automatiza ahora.",
       size=14, bold=True, color=AMBAR, first=True)

    # ---- 7. Nada se pierde
    s = slide_base(prs, "Nada de esto se pierde: todo sirve para Odoo",
                   "Cada avance del puente es también un insumo directo de la implementación")
    filas = [
        ("Maestro único de socios y pólizas", "Es la carga inicial de Contactos y Suscripciones"),
        ("Textos de cobro y recordatorio aprobados", "Son las plantillas del seguimiento automático"),
        ("Política de morosidad validada", "Se configura tal cual como reglas paramétricas"),
        ("Validador SII y su proxy", "Especificación de la validación de RUT (ya compartida con Zoho/soporte)"),
        ("Reglas de negocio verificadas (UF, cuentas, tramos)", "Configuración contable de Odoo sin re-levantamiento"),
        ("Datos limpiados por las validaciones del puente", "Migración más corta y barata"),
    ]
    y = Inches(1.6)
    for hoy, odoo in filas:
        _rect(s, Inches(0.5), y, Inches(5.9), Inches(0.72), GRIS_CLARO)
        tb = _caja(s, Inches(0.62), y + Inches(0.07), Inches(5.7), Inches(0.6))
        _p(tb.text_frame, hoy, size=13, bold=True, first=True, space_after=0)
        tb = _caja(s, Inches(6.5), y + Inches(0.07), Inches(0.5), Inches(0.6))
        _p(tb.text_frame, "→", size=18, bold=True, color=AZUL, first=True, align=PP_ALIGN.CENTER)
        _rect(s, Inches(7.0), y, Inches(5.83), Inches(0.72), AZUL_CLARO)
        tb = _caja(s, Inches(7.12), y + Inches(0.07), Inches(5.6), Inches(0.6))
        _p(tb.text_frame, odoo, size=13, first=True, space_after=0)
        y += Inches(0.82)

    # ---- 8. Trabajo con el equipo MS ampliado
    s = slide_base(prs, "Trabajo con el equipo MS ampliado",
                   "Motivar y evangelizar las virtudes de tener un único sistema")
    bloque(s, Inches(0.5), Inches(1.5), Inches(6.0), Inches(4.9),
           "El mensaje", [
               "Una sola ficha del socio: la misma información para atención, recaudación y finanzas",
               "Cero doble digitación: lo que se escribe una vez, no se vuelve a escribir",
               "El cobro y el seguimiento corren solos; las personas se dedican a los socios",
               "La información del día disponible para decidir, sin esperar el cierre de mes",
           ])
    bloque(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(4.9),
           "Cómo se hace", [
               "Sesiones con el equipo ampliado (finanzas, recaudación, atención, comunicaciones), no solo con los jefes",
               "Las herramientas del puente como demostración: la transformación reduce carga, no la aumenta",
               "Patricio como referente técnico del frente financiero; Marcos ya integrado a recaudación",
               "Alineamiento con la Gerencia General en expectativas, alcance y beneficios",
           ])

    # ---- 9. Decisiones y próximos pasos
    s = slide_base(prs, "Decisiones y próximos pasos")
    bloque(s, Inches(0.5), Inches(1.5), Inches(6.0), Inches(4.9),
           "Decisiones de gerencia", [
               "Partner implementador y modalidad de hosting (condiciona el TXT bancario)",
               "Pasarela de pagos para PAC/PAT y Webpay",
               "Política de morosidad y pagos parciales",
               "Rol de Addval después de la migración",
               "Servicio de consulta SII (~$10.000/mes)",
           ])
    bloque(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(4.9),
           "Próximas semanas", [
               "Entrevistas de cierre y consolidación del maestro único",
               "Agosto en paralelo: herramientas junto al proceso manual",
               "Sesiones con el equipo MS ampliado",
               "Respuestas de los partners a las preguntas críticas → decisión",
               "Kick-off de Odoo con el equipo preparado y los datos limpios",
           ])

    prs.save(salida)
    print(f"OK: {salida} ({salida.stat().st_size:,} bytes)")


if __name__ == "__main__":
    destino = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("Presentacion_Ejecutiva_MS_JR_V1.2.pptx")
    construir(destino)
