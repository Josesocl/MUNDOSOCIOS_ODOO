"""Tests del cruzador de pagos. Ejecutar desde cruzador-pagos/:
    python3 -m unittest discover tests
Las fixtures replican los formatos reales levantados el 2026-07-07
(cartola Banco de Chile como HTML disfrazado de .xls, resumen Transbank).
"""

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import cruzador_pagos as cp

CARTOLA_HTML = """<html><body>
<table>
<tr><td>Sr(a).: Oriana Camila Romero Guerra</td></tr>
<tr><td>Cuenta N°: 00-800-11043-09</td></tr>
<tr><th>Fecha</th><th>Descripción</th><th>Canal o Sucursal</th><th>Nro. Docto.</th>
<th>Cargos (CLP)</th><th>Abonos (CLP)</th><th>Saldo (CLP)</th></tr>
<tr><td>30/06/2026</td><td>Pago: Abonos Debito Y Credito Transbank 0966893109</td>
<td>Oficina Central</td><td></td><td></td><td>989,713</td><td>147,113,606</td></tr>
<tr><td>30/06/2026</td><td>Traspaso De: Enrique Gonzalo Rodriguez Lagos</td>
<td>Internet</td><td></td><td></td><td>40,820</td><td>147,675,986</td></tr>
<tr><td>26/06/2026</td><td>Dep.cheq.otros Bancos</td>
<td>Apoquindo</td><td>8783081</td><td></td><td>700,155</td><td>145,234,184</td></tr>
<tr><td>17/06/2026</td><td>Pac Multibanco Banco Santander</td>
<td>Oficina Central</td><td></td><td></td><td>5,226,809</td><td>211,050,948</td></tr>
<tr><td>25/06/2026</td><td>Provision: Proveedores 00110</td>
<td>Oficina Central</td><td></td><td>60,541,279</td><td></td><td>201,616,120</td></tr>
<tr><td>22/06/2026</td><td>Pago: Abonos Debito Y Credito Transbank 0966893109</td>
<td>Oficina Central</td><td></td><td></td><td>1,014,064</td><td>200,000,000</td></tr>
<tr><th>Fecha</th><th>Descripción</th><th>Canal o Sucursal</th><th>Nro. Docto.</th>
<th>Cargos (CLP)</th><th>Abonos (CLP)</th><th>Saldo (CLP)</th></tr>
<tr><td>Infórmese sobre la garantía estatal de los depósitos</td></tr>
</table></body></html>"""

RESUMEN_HTML = """<html><body><table>
<tr><td>Reporte: Abono acumulado</td></tr>
<tr><th>Fecha de abono</th><th>Total ventas (+)</th><th>Comisión</th><th>Anuladas</th>
<th>Cobros</th><th>Descontados</th><th>Devolución</th><th>Total abono</th>
<th>Cuenta de depósito</th><th>N° de ventas</th></tr>
<tr><td>30/06/2026</td><td>$1,001,977</td><td>$12,264</td><td>$0</td><td>$0</td>
<td>$12,264</td><td>$0</td><td>$989,713</td><td>BANCO DE CHILE 8001104309</td><td>15</td></tr>
<tr><td>22/06/2026</td><td>$1,030,000</td><td>$15,937</td><td>$0</td><td>$0</td>
<td>$15,937</td><td>$0</td><td>$1,014,063</td><td>BANCO DE CHILE 8001104309</td><td>28</td></tr>
<tr><td>05/06/2026</td><td>$0</td><td>$0</td><td>$0</td><td>$0</td><td>$0</td><td>$0</td>
<td>$0</td><td>-</td><td>0</td></tr>
</table></body></html>"""

MAESTRO_CSV = (
    "rut;nombre\n"
    "5031324-7;ENRIQUE GONZALO RODRIGUEZ LAGOS\n"
    "4461179-1;MARIA SOLEDAD PEREZ GARCIA\n"
)


class TestParsers(unittest.TestCase):
    def test_parse_monto(self):
        self.assertEqual(cp.parse_monto("26,025,126"), 26025126)
        self.assertEqual(cp.parse_monto("$1.021.649"), 1021649)
        self.assertEqual(cp.parse_monto("$989,713"), 989713)
        self.assertEqual(cp.parse_monto(""), 0)
        self.assertEqual(cp.parse_monto("-"), 0)
        self.assertEqual(cp.parse_monto(989713), 989713)

    def test_parse_fecha_formatos_mixtos(self):
        self.assertEqual(cp.parse_fecha("30/06/2026"), date(2026, 6, 30))
        self.assertEqual(cp.parse_fecha("6/30/2026"), date(2026, 6, 30))  # re-tipeada
        self.assertEqual(cp.parse_fecha("05/06/2026"), date(2026, 6, 5))  # dd/mm default
        self.assertIsNone(cp.parse_fecha("Fecha"))
        self.assertEqual(cp.parse_fecha(date(2026, 6, 1)), date(2026, 6, 1))

    def test_leer_cartola_html(self):
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "cartola.xls"
            ruta.write_text(CARTOLA_HTML, encoding="utf-8")
            movs = cp.leer_cartola(ruta)
        self.assertEqual(len(movs), 6)  # ignora metadata, headers repetidos y pie
        self.assertEqual(movs[0]["abono"], 989713)
        self.assertEqual(movs[2]["docto"], "8783081")
        self.assertEqual(movs[4]["cargo"], 60541279)

    def test_leer_resumen(self):
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "resumen.xls"
            ruta.write_text(RESUMEN_HTML, encoding="utf-8")
            abonos = cp.leer_resumen_transbank(ruta)
        self.assertEqual(len(abonos), 2)  # descarta el día en $0
        self.assertEqual(abonos[0]["total"], 989713)
        self.assertEqual(abonos[0]["n_ventas"], 15)


class TestClasificacion(unittest.TestCase):
    def _correr(self, tmp):
        tmp = Path(tmp)
        (tmp / "cartola.xls").write_text(CARTOLA_HTML, encoding="utf-8")
        (tmp / "resumen.xls").write_text(RESUMEN_HTML, encoding="utf-8")
        (tmp / "maestro.csv").write_text(MAESTRO_CSV, encoding="utf-8")
        rc = cp.main(["--cartola", str(tmp / "cartola.xls"),
                      "--transbank-resumen", str(tmp / "resumen.xls"),
                      "--maestro", str(tmp / "maestro.csv"),
                      "--salida", str(tmp / "out")])
        self.assertEqual(rc, 0)
        from openpyxl import load_workbook
        ws = load_workbook(tmp / "out" / "PRECONCILIACION_BORRADOR.xlsx").active
        filas = list(ws.iter_rows(min_row=2, values_only=True))
        return {(f[2].day, f[3]): f for f in filas}  # por (día, descripción)

    def test_flujo_completo(self):
        with tempfile.TemporaryDirectory() as d:
            por_desc = self._correr(d)
        transbank = por_desc[(30, "Pago: Abonos Debito Y Credito Transbank 0966893109")]
        self.assertIn("15 ventas", transbank[8])          # CONCEPTO
        self.assertEqual(transbank[14], "listo")          # ESTADO
        traspaso = por_desc[(30, "Traspaso De: Enrique Gonzalo Rodriguez Lagos")]
        self.assertEqual(traspaso[7], "5031324-7")        # RUT propuesto
        self.assertEqual(traspaso[16], "alta")            # CONFIANZA
        pac = por_desc[(17, "Pac Multibanco Banco Santander")]
        self.assertIn("POR DISTRIBUIR", pac[8])
        cheque = por_desc[(26, "Dep.cheq.otros Bancos")]
        self.assertIn("POR IDENTIFICAR", cheque[8])
        prov = por_desc[(25, "Provision: Proveedores 00110")]
        self.assertEqual(prov[15], "CARGO PROVEEDORES")

    def test_match_transbank_tolerancia(self):
        # cartola 22/06 $1,014,064 vs resumen $1,014,063: ±1 peso → match
        with tempfile.TemporaryDirectory() as d:
            por_desc = self._correr(d)
        filas = [f for desc, f in por_desc.items() if "Transbank" in str(desc)]
        estados = {f[14] for f in filas}
        self.assertEqual(estados, {"listo"})

    def test_match_socio_difuso(self):
        socios = [{"rut": "5031324-7", "nombre": "ENRIQUE GONZALO RODRIGUEZ LAGOS",
                   "tokens": cp.tokens_nombre("ENRIQUE GONZALO RODRIGUEZ LAGOS")}]
        socio, score = cp.match_socio("Enrique Gonzalo Rodriguez Lagos", socios)
        self.assertEqual(socio["rut"], "5031324-7")
        self.assertGreaterEqual(score, 0.75)
        # nombre corrupto típico de cartola
        socio, score = cp.match_socio("Enrique Rodriguez Lagos", socios)
        self.assertEqual(socio["rut"], "5031324-7")
        self.assertGreaterEqual(score, 0.5)
        socio, score = cp.match_socio("Juan Perez Soto", socios)
        self.assertLess(score, 0.5)


if __name__ == "__main__":
    unittest.main()


class TestCartolaConvertidaExcel(unittest.TestCase):
    """Regresión: cartola guardada como .xlsx desde Excel (layout real de
    CARTOLA BANCO CHILE JUNIO 26.xlsx, 2026-08-01): columna A vacía,
    preámbulo del banco, encabezado con huecos por celdas combinadas y
    montos como texto."""

    def test_lee_cartola_convertida(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "CARTOLA BANCO CHILE JUNIO 26.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Hoja1"
            filas = [
                [None, "Sr(a).: ", None, "Oriana Camila Romero Guerra"],
                [None, "Cuenta N°:", None, "00-800-11043-09"],
                [None, "Total Cargos", None, None, "Total Abonos", None,
                 None, "Línea de Sobregiro Pactado", "Línea de Crédito"],
                [None, "309288233", None, None, "279322268", None, None,
                 "0", "0"],
                [None, "Movimientos", None, "al 07/07/2026"],
                [None, "Fecha", None, "Descripción", None,
                 "Canal o Sucursal", "Nro. Docto.", "Cargos (CLP)",
                 "Abonos (CLP)", "Saldo (CLP)"],
                [None, "30/06/2026", None, "Traspaso De: Enrique Gonza",
                 None, "Internet", None, None, "40820", "147675986"],
                [None, "30/06/2026", None, "Pago: Abonos Debito Y Cred",
                 None, "Oficina Central", None, None, "989713", "147113606"],
                [None, "26/06/2026", None, "Dep.cheq.otros Bancos", None,
                 "Apoquindo", "8783081", None, "700155", "145234184"],
                # encabezado repetido de la página siguiente + fila legal
                [None, "Fecha", None, "Descripción", None,
                 "Canal o Sucursal", "Nro. Docto.", "Cargos (CLP)",
                 "Abonos (CLP)", "Saldo (CLP)"],
                [None, "25/06/2026", None, "Pac Falabella", None,
                 "Oficina Central", None, "1234", None, "145000000"],
                [None, "Información legal del banco…"],
            ]
            for f in filas:
                ws.append(f)
            wb.save(ruta)

            movs = cp.leer_cartola(ruta)
        self.assertEqual(len(movs), 4)
        self.assertEqual(movs[0]["fecha"], date(2026, 6, 30))
        self.assertEqual(movs[0]["abono"], 40820)
        self.assertEqual(movs[0]["descripcion"], "Traspaso De: Enrique Gonza")
        self.assertEqual(movs[2]["docto"], "8783081")
        self.assertEqual(movs[2]["canal"], "Apoquindo")
        # fila de la página 2, con cargo en vez de abono
        self.assertEqual(movs[3]["cargo"], 1234)
        # las filas de saldos/preambulo no entran como movimientos
        montos = [m["abono"] for m in movs]
        self.assertNotIn(279322268, montos)


class TestResumenTransbankConvertido(unittest.TestCase):
    """Regresión: Resumen_historico_abonos guardado como .xlsx desde Excel
    (layout real 2026-08-01): columna A vacía, bloque de totales arriba,
    encabezado duplicado en las filas 29-30 y montos como texto."""

    def test_lee_resumen_convertido(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "Resumen_historico_abonos (06-30).xlsx"
            wb = Workbook()
            ws = wb.active
            filas = [
                [None, "Reporte:", "Abono acumulado"],
                [None, "Total ventas (+)", None, "35700751"],
                [None, "Total abono", None, "34931862"],   # bloque de totales
                [None, "Fecha de abono", "Total ventas (+)",
                 "Montos Descontados", None, None, None,
                 "Devolución comisión po", "Total abono",
                 "Cuenta de depósito", "N° de ventas"],
                [None, "Fecha de abono", "Total ventas (+)",
                 "Comisión Transbank + I", "Ventas Anuladas (-)",
                 "Cobros por servicio** ", "Total montos descontad",
                 "Devolución comisión po", "Total abono",
                 "Cuenta de depósito", "N° de ventas"],
                [None, "30/06/2026", "1001977", "12264", "0", "0", "12264",
                 "0", "989713", "BANCO DE CHILE 8001104", "15"],
                [None, "29/06/2026", "0", "0", "0", "0", "0", "0", "0",
                 "-", "0"],
                [None, "25/06/2026", "444067", "4877", "24000", "0", "28877",
                 "288", "415478", "BANCO DE CHILE 8001104", "7"],
            ]
            for f in filas:
                ws.append(f)
            wb.save(ruta)
            abonos = cp.leer_resumen_transbank(ruta)

        self.assertEqual(len(abonos), 2)      # los días en 0 no entran
        self.assertEqual(abonos[0]["fecha"], date(2026, 6, 30))
        self.assertEqual(abonos[0]["total"], 989713)
        self.assertEqual(abonos[0]["n_ventas"], 15)
        self.assertEqual(abonos[1]["total"], 415478)
        # el bloque de totales de arriba no debe colarse como un abono
        self.assertNotIn(34931862, [a["total"] for a in abonos])
