"""Tests del consolidador de mantenedores (sin archivos reales).

Los fixtures replican el layout levantado el 2026-07-31 de los archivos
reales: encabezados con columnas VALOR duplicadas por año, nóminas
PAC/PAT pegadas debajo del mantenedor, filas de totales, socios nuevos
sin Mod. Pago, y en cuota social empresa el N° de miembros anotado solo
cuando es > 3.

Ejecutar desde consolidador-maestro/:
    python3 -m unittest discover tests
"""

import csv
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import consolidador_maestro as cm

# RUTs con dígito verificador válido (módulo 11)
R1, R2, R3, R4 = "76138907-6", "76191369-7", "2530068-8", "15845066-6"
R5, R6 = "9528348-9", "8615665-2"
RUT_MALO = "12345678-0"  # DV incorrecto


def _plan_socios(ruta):
    wb = Workbook()
    ws = wb.active
    ws.append(["PLAN SOCIOS 2024", None, None, None, "VALOR UF",
               39728, 40820])
    ws.append(["k", "Mod. Pago", "RUT", "NOMBRE", "VALOR '25",
               "VALOR UF actual", "PAC", "PAT"])
    ws.append([2, "DIRECTA", R5, "FRANCISCO JAVIER ROMO CONCHA", 3.55, 3, None, None])
    ws.append([3, "PAC", R6, "CARLOS RODRIGO SEGUEL HINTZ", 2, 2.43, 170, None])
    ws.append([4, "PAT", RUT_MALO, "RUT CON DV MALO", 1, 1, None, 1053])
    ws.append([None, None, R1, "SOCIO NUEVO SIN MODO", None, 2.5, None, None])
    ws.append([None, None, None, None, None, 28428164, None, None])  # total
    # nómina PAC pegada debajo (no debe consolidarse)
    ws.append(["RUT", "ID", "NOMBRE SEGURO", "Monto", "Monto",
               "Fecha en que se efectuó Cargo", "Banco", "Estado de Cargo",
               "Convenio"])
    ws.append([R2, "761913697", "OTRO SOCIO P SOC", 0, 99000,
               "6/16/2026", "CHILE", "cargo efectuado", 16])
    wb.save(ruta)


def _carreno(ruta):
    wb = Workbook()
    ws = wb.active
    ws.append(["P. CARREÑO 2024", None, None, None, "VALOR UF", 40845])
    ws.append(["nº", "Mod. Pago", "Rut", "NOMBRE", "INICIO COBERTURA",
               "VALOR", "ESTADO PAC", "ESTADO PAT"])
    ws.append([2, "DIRECTA", R3, "Norman Goijberg Rein", "10-01-2017",
               0.4, None, None])
    ws.append([3, "PAT", R4, "Pablo Irarrazaval Barros", "10-01-2017",
               "0,2", None, "ACTIVO"])
    wb.save(ruta)


def _cs_empresa(ruta):
    wb = Workbook()
    ws = wb.active
    ws.append(["Nombre de Socio", "RUT", "Número de Socio", "Tipo Socio",
               "Cámara_Regional", None, None, None, "GLOSA"])
    # base 3 UF ($119.195), sin miembros anotados
    ws.append(["Constructora Paolo Brizzi", R1, 91827, "Empresa",
               "O'HIGGINS", 119195, None, 119195,
               f"{R1} CE 26 O'HIGGINS"])
    # 4 miembros anotados
    ws.append(["Constructora RGO Spa", R2, 92067, "Empresa",
               "CONCEPCIÓN", 158927, 4, 158927,
               f"{R2} CE 26 CONCEPCIÓN"])
    # 5 miembros SIN anotar => derivar del monto (119195 + 2*39732)
    ws.append(["Empresa Derivada", R5, 92100, "Empresa", "SANTIAGO",
               198659, None, 198659, f"{R5} CE 26 SANTIAGO"])
    wb.save(ruta)


def _cs_persona(ruta):
    wb = Workbook()
    ws = wb.active
    ws.append(["Nombre de Socio", "RUT", "Número de Socio", "Tipo Socio",
               "Cámara_Regional", "PERSONA", "GRUPO ALERCE", "GLOSA"])
    ws.append(["Carlos Daetz Hofmann", R3, 92160, "Persona", "VALDIVIA",
               39732, None, f"{R3} CP 26 VALDIVIA"])
    ws.append(["Alvaro Sailer Lantadilla", R4, 92025, "Persona", "SANTIAGO",
               39732, None, f"{R4} CP 26 SANTIAGO"])
    # excepción de monto
    ws.append(["Socio Excepcion", R6, 92001, "Persona", "SANTIAGO",
               39587, None, f"{R6} CP 26 SANTIAGO"])
    wb.save(ruta)


class TestConsolidador(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.carpeta = Path(self.tmp.name) / "RECAUDACION"
        self.salida = Path(self.tmp.name) / "salida"
        self.carpeta.mkdir()
        _plan_socios(self.carpeta / "MANT. PLAN SOCIOS 07-26.xlsx")
        _carreno(self.carpeta / "MANT. PLAN CARREÑO 07-26.xlsx")
        _cs_empresa(self.carpeta / "Mantenedor Cuota Social empresa.xlsx")
        _cs_persona(self.carpeta / "Mantenedor Cuota Social persona.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def _correr(self):
        rc = cm.main(["--carpeta", str(self.carpeta),
                      "--salida", str(self.salida)])
        self.assertEqual(rc, 0)

    def test_consolida_y_corta_en_la_nomina_pac(self):
        self._correr()
        with open(self.salida / "maestro_seguros.csv", encoding="utf-8-sig") as f:
            filas = list(csv.DictReader(f, delimiter=";"))
        ruts_ps = [f["rut"] for f in filas if f["seguro"] == "PLAN SOCIOS"]
        # R2 solo aparece en la nómina PAC pegada: NO debe estar
        self.assertNotIn(R2, ruts_ps)
        # el RUT inválido no pasa al CSV
        self.assertNotIn(RUT_MALO, ruts_ps)
        # socio nuevo sin Mod. Pago entra igual, con medio_pago vacío
        # (el generador lo acepta: el campo es opcional/informativo)
        nuevo = next(f for f in filas if f["rut"] == R1)
        self.assertEqual(nuevo["medio_pago"], "")
        self.assertEqual(nuevo["factor_uf"], "2,5")

    def test_factor_es_la_columna_valor_mas_a_la_derecha(self):
        self._correr()
        with open(self.salida / "maestro_seguros.csv", encoding="utf-8-sig") as f:
            filas = {(r["rut"], r["seguro"]): r
                     for r in csv.DictReader(f, delimiter=";")}
        # Plan Socios: VALOR UF actual (3), no VALOR '25 (3.55)
        self.assertEqual(filas[(R5, "PLAN SOCIOS")]["factor_uf"], "3,0")
        # Carreño: factor '0,2' en texto se lee bien
        self.assertEqual(filas[(R4, "PLAN CARRENO")]["factor_uf"], "0,2")

    def test_cuota_social_miembros_y_excepciones(self):
        self._correr()
        with open(self.salida / "maestro_cuota_social.csv",
                  encoding="utf-8-sig") as f:
            filas = {r["rut"]: r for r in csv.DictReader(f, delimiter=";")}
        self.assertEqual(filas[R1]["miembros"], "3")     # base
        self.assertEqual(filas[R2]["miembros"], "4")     # anotado
        self.assertEqual(filas[R5]["miembros"], "5")     # derivado del monto
        self.assertEqual(filas[R6]["monto_clp"], "39587")  # excepción persona
        self.assertEqual(filas[R3]["monto_clp"], "")     # normal, sin override
        self.assertEqual(filas[R3]["camara"], "VALDIVIA")

    def test_maestro_excel_incluye_observados_y_resumen(self):
        self._correr()
        wb = load_workbook(self.salida / "MAESTRO_UNICO_MS.xlsx")
        filas = list(wb["MAESTRO"].iter_rows(min_row=2, values_only=True))
        ruts = [f[0] for f in filas]
        self.assertIn(RUT_MALO, ruts)   # el Excel sí lo muestra, con observación
        fila_mala = next(f for f in filas if f[0] == RUT_MALO)
        self.assertIn("RUT INVALIDO", fila_mala[10])
        resumen = {r[0]: r[1] for r in
                   wb["RESUMEN"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(resumen["PLAN SOCIOS"], 4)
        self.assertEqual(resumen["CUOTA SOCIAL EMPRESA"], 3)

    def test_avisa_mantenedores_faltantes(self):
        # sin COMPLEMENTARIO ni CATASTROFICO en la carpeta: corre igual
        self._correr()
        self.assertTrue((self.salida / "MAESTRO_UNICO_MS.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
